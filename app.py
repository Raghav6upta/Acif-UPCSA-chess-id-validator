import io
import re
import time
import threading
import webbrowser
from datetime import datetime
from pathlib import Path

import requests
from bs4 import BeautifulSoup
from flask import Flask, jsonify, render_template, request, send_file
from openpyxl import load_workbook


# =========================================================
# APP SETUP
# =========================================================

APP_DIR = Path(__file__).resolve().parent

app = Flask(
    __name__,
    template_folder=str(APP_DIR / "templates")
)


# =========================================================
# URLS
# =========================================================

AICF_URL = "https://admin.aicf.in/api/players"
UPCSA_URL = "https://www.upchess.org/view_player.php"


# =========================================================
# HTTP SESSION
# =========================================================

session = requests.Session()

session.headers.update({
    "User-Agent": (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
        "AppleWebKit/537.36 (KHTML, like Gecko) "
        "Chrome/131.0.0.0 Safari/537.36"
    )
})


# =========================================================
# GENERAL HELPER
# =========================================================

def clean(value):
    if value is None:
        return ""

    return re.sub(r"\s+", " ", str(value)).strip()


# =========================================================
# TABLE FIELD HELPER
# =========================================================

def text_after_label(soup, label):

    for cell in soup.find_all(["td", "th"]):

        cell_text = clean(
            cell.get_text(" ", strip=True)
        )

        if cell_text.lower() == label.lower():

            next_cell = cell.find_next_sibling(
                ["td", "th"]
            )

            if next_cell:
                return clean(
                    next_cell.get_text(
                        " ",
                        strip=True
                    )
                )

    return ""


# =========================================================
# AICF LOOKUP
# =========================================================

def aicf_lookup(player_id):

    r = session.get(
        AICF_URL,
        params={
            "name": player_id,
            "state": 0,
            "city": 0
        },
        timeout=20
    )

    r.raise_for_status()

    payload = r.json()

    results = payload.get(
        "data",
        []
    )

    # AICF may return multiple search results.
    # We therefore require an exact ID match.

    player = next(
        (
            item
            for item in results
            if clean(
                item.get("aicf_id")
            ) == player_id
        ),
        None
    )

    # -----------------------------------------------------
    # PLAYER NOT FOUND
    # -----------------------------------------------------

    if not player:

        return {
            "verification_status": "Not Found",
            "aicf_id": "",
            "name": "",
            "gender": "",
            "city": "",
            "district": "",
            "state": "",
            "membership_status": "",
            "membership_expiry": ""
        }

    # -----------------------------------------------------
    # MEMBERSHIP STATUS
    # -----------------------------------------------------

    status = (
        "Active"
        if player.get("membership_status")
        else "Not Active"
    )

    # -----------------------------------------------------
    # MEMBERSHIP EXPIRY
    # -----------------------------------------------------

    expiry = clean(
        player.get("membership_expire_at")
    )

    if expiry:

        try:

            expiry = datetime.fromisoformat(
                expiry.replace(
                    "Z",
                    "+00:00"
                )
            ).strftime(
                "%d-%b-%Y"
            )

        except Exception:
            pass

    # -----------------------------------------------------
    # NAME
    # -----------------------------------------------------

    name = " ".join(
        part
        for part in [
            clean(
                player.get("first_name")
            ),
            clean(
                player.get("middle_name")
            ),
            clean(
                player.get("last_name")
            )
        ]
        if part
    )

    # -----------------------------------------------------
    # RETURN
    # -----------------------------------------------------

    return {
        "verification_status": "Found",

        "aicf_id": clean(
            player.get("aicf_id")
        ),

        "name": name,

        "gender": clean(
            player.get("gender")
        ),

        "city": clean(
            player.get("city_name")
        ),

        "district": clean(
            player.get("district_name")
        ),

        "state": clean(
            player.get("state_name")
        ),

        "membership_status": status,

        "membership_expiry": expiry
    }


# =========================================================
# UPCSA LOOKUP
# =========================================================

def upcsa_lookup(player_id):

    # -----------------------------------------------------
    # BROWSER-LIKE HEADERS
    # -----------------------------------------------------

    headers = {
        "User-Agent": (
            "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
            "AppleWebKit/537.36 (KHTML, like Gecko) "
            "Chrome/131.0.0.0 Safari/537.36"
        ),

        "Referer":
            "https://www.upchess.org/",

        "Accept": (
            "text/html,application/xhtml+xml,"
            "application/xml;q=0.9,"
            "*/*;q=0.8"
        ),

        "Accept-Language":
            "en-US,en;q=0.9",

        "Connection":
            "keep-alive"
    }

    # -----------------------------------------------------
    # RATE LIMIT
    #
    # Wait 4 seconds before every UPCSA request.
    # -----------------------------------------------------

    time.sleep(4)

    last_error = None
    response = None

    # -----------------------------------------------------
    # TRY UP TO 3 TIMES
    # -----------------------------------------------------

    for attempt in range(3):

        try:

            response = session.get(
                UPCSA_URL,
                params={
                    "id": player_id
                },
                headers=headers,
                timeout=30
            )

            response.raise_for_status()

            # Successful request
            break

        except requests.RequestException as exc:

            last_error = exc

            # Retry if this wasn't the final attempt.
            if attempt < 2:

                if response is not None:

                    # Retry delay:
                    #
                    # attempt 1 -> 8 seconds
                    # attempt 2 -> 15 seconds

                    delay = (
                        8
                        if attempt == 0
                        else 15
                    )

                else:
                    delay = 8

                time.sleep(delay)

    else:

        status_code = (
            response.status_code
            if response is not None
            else "unknown"
        )

        raise requests.RequestException(
            f"UPCSA request failed "
            f"(HTTP {status_code}) "
            f"for player ID {player_id} "
            f"after 3 attempts."
        ) from last_error

    # =====================================================
    # PARSE HTML
    # =====================================================

    soup = BeautifulSoup(
        response.text,
        "html.parser"
    )

    page_text = clean(
        soup.get_text(
            " ",
            strip=True
        )
    )

    # =====================================================
    # BASIC PLAYER INFORMATION
    # =====================================================

    data = {

        "up_player_id":
            player_id,

        "name":
            text_after_label(
                soup,
                "Name"
            ),

        "email":
            text_after_label(
                soup,
                "Email"
            ),

        "mobile":
            text_after_label(
                soup,
                "Mobile No"
            ),

        "father_name":
            text_after_label(
                soup,
                "Father's Name"
            ),

        "gender":
            text_after_label(
                soup,
                "Gender"
            ),

        "year_of_birth":
            text_after_label(
                soup,
                "Year of Birth"
            ),

        "address":
            text_after_label(
                soup,
                "Address"
            ),

        "aicf_id":
            text_after_label(
                soup,
                "AICF ID"
            ),

        "fide_id":
            text_after_label(
                soup,
                "FIDE ID"
            ),

        "registered_as":
            text_after_label(
                soup,
                "Registered As:"
            ),

        "district":
            "",

        "state":
            "",

        "membership_status":
            "Unknown",

        "membership_expiry":
            ""
    }

    # =====================================================
    # DISTRICT + STATE
    # =====================================================

    for cell in soup.find_all(
        ["td", "th"]
    ):

        text = clean(
            cell.get_text(
                " ",
                strip=True
            )
        )

        if (
            "District:" in text
            and
            "State:" in text
        ):

            match = re.search(
                r"District:\s*(.*?)\s+State:\s*(.*)$",
                text,
                re.IGNORECASE
            )

            if match:

                data["district"] = clean(
                    match.group(1)
                )

                data["state"] = clean(
                    match.group(2)
                )

                break

    # =====================================================
    # MEMBERSHIP STATUS
    # =====================================================

    # UPCSA inactive page contains:
    #
    # Not Active! Renew Your Membership.
    #

    inactive_match = re.search(
        r"Not\s*Active!\s*Renew\s*Your\s*Membership",
        page_text,
        re.IGNORECASE
    )

    # UPCSA active page contains something like:
    #
    # Valid Up To : 31 March 2027
    #

    active_match = re.search(
        r"Valid\s*Up\s*To\s*:?\s*"
        r"([0-9]{1,2}\s+[A-Za-z]+\s+[0-9]{4})",
        page_text,
        re.IGNORECASE
    )

    if inactive_match:

        data["membership_status"] = "Not Active"
        data["membership_expiry"] = ""

    elif active_match:

        data["membership_status"] = "Active"

        data["membership_expiry"] = clean(
            active_match.group(1)
        )

    else:

        data["membership_status"] = "Unknown"
        data["membership_expiry"] = ""

    # =====================================================
    # VERIFICATION STATUS
    # =====================================================

    if data["name"]:

        data["verification_status"] = "Found"

    else:

        data["verification_status"] = "Not Found"

    return data


# =========================================================
# HOME PAGE
# =========================================================

@app.get("/")
def index():

    return render_template(
        "index.html"
    )


# =========================================================
# PROCESS EXCEL
# =========================================================

@app.post("/process")
def process():

    file = request.files.get(
        "file"
    )

    mode = request.form.get(
        "mode",
        ""
    ).lower()

    # -----------------------------------------------------
    # VALIDATE FILE
    # -----------------------------------------------------

    if not file:

        return jsonify({
            "error":
                "No Excel file uploaded."
        }), 400

    # -----------------------------------------------------
    # VALIDATE MODE
    # -----------------------------------------------------

    if mode not in {
        "aicf",
        "upcsa"
    }:

        return jsonify({
            "error":
                "Invalid mode."
        }), 400

    # -----------------------------------------------------
    # READ EXCEL
    # -----------------------------------------------------

    original = file.read()

    wb = load_workbook(
        io.BytesIO(original)
    )

    ws = wb.active

    # -----------------------------------------------------
    # READ HEADERS
    # -----------------------------------------------------

    headers = [
        clean(cell.value)
        for cell in ws[1]
    ]

    # -----------------------------------------------------
    # FIND PLAYER ID COLUMN
    # -----------------------------------------------------

    id_idx = None

    for i, header in enumerate(
        headers,
        start=1
    ):

        if header.lower() in {
            "playersid",
            "player id",
            "playerid",
            "aicf id",
            "up player id"
        }:

            id_idx = i

            break

    if id_idx is None:

        return jsonify({
            "error": (
                "Could not find a "
                "PlayersID / Player ID "
                "column in row 1."
            )
        }), 400

    # =====================================================
    # OUTPUT COLUMNS
    # =====================================================

    if mode == "aicf":

        columns = [

            "AICF Verification Status",

            "AICF ID",

            "AICF Name",

            "AICF Gender",

            "AICF City",

            "AICF District",

            "AICF State",

            "AICF Membership Status",

            "AICF Membership Expiry"
        ]

        lookup = aicf_lookup

    else:

        columns = [

            "UPCSA Verification Status",

            "UPCSA Name",

            "UPCSA Father's Name",

            "UPCSA Gender",

            "UPCSA Year of Birth",

            "UPCSA District",

            "UPCSA State",

            "UPCSA Registered As",

            "UPCSA AICF ID",

            "UPCSA FIDE ID",

            "UPCSA Membership Status",

            "UPCSA Membership Expiry",

            "UPCSA Email (masked)",

            "UPCSA Mobile (masked)",

            "UPCSA Address (masked)"
        ]

        lookup = upcsa_lookup

    # =====================================================
    # ADD OUTPUT HEADERS
    # =====================================================

    start_col = (
        ws.max_column + 1
    )

    for offset, column in enumerate(
        columns
    ):

        ws.cell(
            row=1,
            column=start_col + offset,
            value=column
        )

    # =====================================================
    # CACHE
    # =====================================================

    # If the same player ID appears multiple times,
    # only make one request.

    cache = {}

    # =====================================================
    # PROCESS PLAYERS
    # =====================================================

    for row in range(
        2,
        ws.max_row + 1
    ):

        player_id = clean(
            ws.cell(
                row=row,
                column=id_idx
            ).value
        )

        if not player_id:
            continue

        # -------------------------------------------------
        # LOOKUP
        # -------------------------------------------------

        if player_id not in cache:

            try:

                cache[player_id] = lookup(
                    player_id
                )

            except Exception as exc:

                cache[player_id] = {

                    "verification_status":
                        f"Error: {type(exc).__name__}",

                    "error":
                        str(exc)
                }

        data = cache[player_id]

        # =================================================
        # AICF RESULTS
        # =================================================

        if mode == "aicf":

            values = [

                data.get(
                    "verification_status",
                    ""
                ),

                data.get(
                    "aicf_id",
                    ""
                ),

                data.get(
                    "name",
                    ""
                ),

                data.get(
                    "gender",
                    ""
                ),

                data.get(
                    "city",
                    ""
                ),

                data.get(
                    "district",
                    ""
                ),

                data.get(
                    "state",
                    ""
                ),

                data.get(
                    "membership_status",
                    ""
                ),

                data.get(
                    "membership_expiry",
                    ""
                )
            ]

        # =================================================
        # UPCSA RESULTS
        # =================================================

        else:

            values = [

                data.get(
                    "verification_status",
                    ""
                ),

                data.get(
                    "name",
                    ""
                ),

                data.get(
                    "father_name",
                    ""
                ),

                data.get(
                    "gender",
                    ""
                ),

                data.get(
                    "year_of_birth",
                    ""
                ),

                data.get(
                    "district",
                    ""
                ),

                data.get(
                    "state",
                    ""
                ),

                data.get(
                    "registered_as",
                    ""
                ),

                data.get(
                    "aicf_id",
                    ""
                ),

                data.get(
                    "fide_id",
                    ""
                ),

                data.get(
                    "membership_status",
                    ""
                ),

                data.get(
                    "membership_expiry",
                    ""
                ),

                data.get(
                    "email",
                    ""
                ),

                data.get(
                    "mobile",
                    ""
                ),

                data.get(
                    "address",
                    ""
                )
            ]

        # =================================================
        # WRITE DATA
        # =================================================

        for offset, value in enumerate(
            values
        ):

            ws.cell(
                row=row,
                column=start_col + offset,
                value=value
            )

    # =====================================================
    # SAVE EXCEL
    # =====================================================

    output = io.BytesIO()

    wb.save(output)

    output.seek(0)

    # -----------------------------------------------------
    # OUTPUT FILE NAME
    # -----------------------------------------------------

    name = Path(
        file.filename or "players.xlsx"
    ).stem

    output_name = (
        f"{name}_verified.xlsx"
    )

    # -----------------------------------------------------
    # DOWNLOAD
    # -----------------------------------------------------

    return send_file(
        output,
        as_attachment=True,
        download_name=output_name,
        mimetype=(
            "application/vnd.openxmlformats-"
            "officedocument.spreadsheetml.sheet"
        )
    )


# =========================================================
# OPEN BROWSER
# =========================================================

def open_browser():

    webbrowser.open(
        "http://127.0.0.1:8765"
    )


# =========================================================
# START SERVER
# =========================================================

if __name__ == "__main__":

    threading.Timer(
        0.8,
        open_browser
    ).start()

    app.run(
        host="127.0.0.1",
        port=8765,
        debug=False
    )